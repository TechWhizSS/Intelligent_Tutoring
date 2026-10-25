"""Build a compact illustrated PDF and editable Word report from real UI captures."""
from pathlib import Path
import json,html
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor,white
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs/report';SHOTS=OUT/'screenshots'
WEB='https://syllabus-to-song-studio.aptcat4096.chatgpt.site'
APK='https://github.com/TechWhizSS/Intelligent_Tutoring/raw/refs/heads/codex/development/android/downloads/SyllabusToSong-v1.0-sunrise.apk'
REPO='https://github.com/TechWhizSS/Intelligent_Tutoring/tree/codex/development'
DATE='25 October 2026'
FONT='/usr/share/fonts/truetype/open-sans/'
pdfmetrics.registerFont(TTFont('OpenSans',FONT+'OpenSans-Regular.ttf'))
pdfmetrics.registerFont(TTFont('OpenSans-Bold',FONT+'OpenSans-Bold.ttf'))
pdfmetrics.registerFontFamily('OpenSans',normal='OpenSans',bold='OpenSans-Bold',italic='OpenSans',boldItalic='OpenSans-Bold')
W,H=A4;M=42;CW=W-2*M
INK='#302421';MUTED='#7C6A63';ACCENT='#C64620';PEACH='#FFF0E6';LINE='#EFDFD4'


def para(c,text,x,top,width,size=9.2,leading=14,color=INK,bold=False):
    style=ParagraphStyle('body',fontName='OpenSans-Bold' if bold else 'OpenSans',fontSize=size,leading=leading,textColor=HexColor(color))
    p=Paragraph(text.replace('→',' &gt; ').replace('↗',''),style);_,height=p.wrap(width,H);p.drawOn(c,x,top-height);return top-height

def box(c,x,y,w,h,color=PEACH):
    c.setFillColor(HexColor(color));c.setStrokeColor(HexColor(LINE));c.roundRect(x,y,w,h,10,stroke=1,fill=1)

def image(c,name,x,top,w,h):
    path=SHOTS/name
    iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale
    c.setFillColor(white);c.setStrokeColor(HexColor(LINE));c.roundRect(x-3,top-h-3,w+6,h+6,6,fill=1,stroke=1)
    c.drawImage(str(path),x+(w-dw)/2,top-dh,width=dw,height=dh,mask='auto')

def base(c,page,label):
    c.setFillColor(HexColor('#FFFCF8'));c.rect(0,0,W,H,fill=1,stroke=0)
    para(c,'SYLLABUS TO SONG  /  '+label,M,H-29,CW,8,12,ACCENT,True)
    c.setStrokeColor(HexColor(LINE));c.line(M,42,W-M,42)
    para(c,'MBA Tech · CE · First year · Semester I',M,32,CW-80,7.3,10,MUTED)
    para(c,f'{DATE}  /  {page}',W-175,32,133,7.3,10,MUTED)


def pdf_report():
    c=canvas.Canvas(str(OUT/'Syllabus_to_Song_Compact_Report.pdf'),pagesize=A4)
    c.setTitle('Syllabus to Song — Website and Android App: Compact Project Report')
    c.setAuthor('Syllabus to Song project')
    base(c,1,'PROJECT REPORT')
    para(c,'One project. Two working interfaces.',M,779,CW,24,30,INK,True)
    para(c,'We built <b>both a responsive full-stack website and a native Android app</b> for music-assisted Semester I learning. The app adapts the uploaded Android project; both share the approved syllabus and reddish-orange sunrise design.',M,733,CW,10,15)
    for i,(number,label) in enumerate([('2','Deliverables'),('91','Uploaded topics'),('10','Ordered units'),('EN / HI','Languages')]):
        w=(CW-24)/4;x=M+i*(w+8);box(c,x,622,w,61)
        para(c,number,x+13,670,w-26,20,24,ACCENT,True);para(c,label,x+13,643,w-26,8,11,MUTED)
    para(c,'Website: live, phone-friendly dashboard',M,605,248,10,15,INK,True)
    para(c,'Android: installable demonstration APK',M+263,605,248,10,15,INK,True)
    para(c,'Sample SAP login, subject/chapter controls, lyrics, study tools, library/player and reviewed material uploads.',M,582,248,8.7,13)
    para(c,'Android 8.0+, local drafts/library, native audio controls, study tools, ACE-Step connection and Gemini handoff.',M+263,582,248,8.7,13)
    para(c,f'<b>Open the webpage:</b> <link href="{WEB}" color="{ACCENT}">{WEB}</link>',M,533,CW,8.7,14)
    para(c,'Website access: owner-private; sign in with the owning ChatGPT account if prompted.',M,516,CW,7.3,10,MUTED)
    image(c,'web-studio.png',M,499,CW,350)
    para(c,'Website dashboard — real current UI with sample lyric drafts; no generated song is shown.',M,137,CW,8,12,MUTED)
    para(c,'Shared learning flow',M,108,CW,10,14,INK,True)
    para(c,'SAP login → MBA Tech / CE / Year 1 / Semester I → Calculus or BEEE → unit/topic → English or Hindi → style and lyrics → music route → permitted audio playback/export.',M,88,CW,9,14)
    c.showPage()

    base(c,2,'ANDROID APP')
    para(c,'The native Android companion',M,779,CW,25,31,INK,True)
    para(c,'Kotlin + Jetpack Compose · SQLite song library · Android MediaPlayer · device-local drafts',M,739,CW,9,14,MUTED)
    para(c,f'<link href="{APK}" color="{ACCENT}"><b>Download the Android APK ↗</b></link> &nbsp; Install it, then tap “Try the sample student account”.',M,713,CW,9,14)
    gap=15;pw=(CW-2*gap)/3
    for i,(name,title,caption) in enumerate([
        ('android-studio.png','Chapter selection','Programme, semester, subject, unit and topic.'),
        ('android-music.png','Lyrics & music controls','Editable drafts and the ACE-Step synthesis action.'),
        ('android-notes.png','Source-linked study','Uploaded summaries and labelled supplemental guidance.')]):
        x=M+i*(pw+gap);image(c,name,x,680,pw,360)
        para(c,title,x,306,pw,9.2,13,INK,True);para(c,caption,x,287,pw,8,12,MUTED)
    para(c,'Screens above are real app UI rendered with Robolectric, not physical-phone screenshots. No audio was generated for these captures.',M,249,CW,7.8,12,MUTED)
    box(c,M,116,CW,107)
    para(c,'What the app adds',M+15,209,CW-30,11,15,ACCENT,True)
    para(c,'• Genre, mood, vocal delivery, tempo, length, difficulty, formula repeats and instructions.<br/>• Local song history/favourites; play, pause, seek, replay, permitted audio import and export.<br/>• Notes, bilingual quizzes, all-topic completeness and sample-admin material review.<br/>• GPU connection-file import with encrypted key storage; Gemini browser fallback.',M+15,184,CW-30,8.8,14)
    para(c,'ACE-Step music needs an owner-run GPU',M,103,CW,10,14,INK,True)
    para(c,'More → Music connection → connect the running free GPU notebook → generate/edit lyrics → Synthesize with ACE-Step. Free GPU availability and session limits apply. The phone does not host the model.',M,82,CW,8.5,13)
    c.showPage()

    base(c,3,'FEATURES & VALIDATION')
    para(c,'Learning features, with clear status',M,779,CW,24,30,INK,True)
    gallery=[('web-notes.png','Chapter notes & formula references'),('web-quiz.png','Supplemental quizzes with feedback'),('web-coverage.png','Completeness report: every uploaded topic'),('web-gemini.png','Website music: Gemini / Lyria handoff')]
    gap=18;gw=(CW-gap)/2
    for i,(name,caption) in enumerate(gallery):
        row,col=divmod(i,2);x=M+col*(gw+gap);top=724-row*202
        image(c,name,x,top,gw,172)
        para(c,caption,x,top-181,gw,8.3,12,INK,True)
    para(c,'Academic scope',M,304,CW,10.5,15,ACCENT,True)
    para(c,'38 Calculus + 53 BEEE topics, in uploaded unit/topic order. Source excerpts, formulas, locators and supplied objectives are preserved. Full chapter notes were not supplied; generated guidance, lyric templates and quizzes are labelled supplemental.',M,282,CW,9,14)
    para(c,'Verified engineering work',M,226,CW,10.5,15,ACCENT,True)
    para(c,'38 Python checks; 7 ACE-Step protocol tests; 2 Robolectric/Compose flow tests; Android build/lint and APK signature verification. Website browser checks cover the song flow, all 91 report rows, fixture audio playback and 320/390/768px layouts. GitHub website and Android workflows passed.',M,204,CW,8.8,14)
    box(c,M,68,CW,79)
    para(c,'Music-generation boundary',M+14,134,CW-28,10,14,ACCENT,True)
    para(c,'<b>No real educational song or singing quality has been verified through this project.</b> Gemini opens externally: choose Create music, paste and send. ACE-Step needs a running GPU and an actual generation/listening test. No song is fabricated. App and website libraries are separate; physical-device audio decoding remains unverified.',M+14,114,CW-28,8.2,12)
    para(c,f'<link href="{REPO}" color="{ACCENT}">Source code: TechWhizSS/Intelligent_Tutoring · codex/development</link>',M,57,CW,7.5,10)
    c.save()


def hyperlink(paragraph,label,url):
    node=OxmlElement('w:hyperlink');node.set(qn('r:id'),paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'C64620');props.append(col);run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);node.append(run);paragraph._p.append(node)


def word_report():
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.55);sec.left_margin=sec.right_margin=Inches(.6)
    normal=d.styles['Normal'];normal.font.name='Open Sans';normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(6)
    for name in ['Title','Heading 1','Heading 2']:
        d.styles[name].font.name='Open Sans';d.styles[name].font.color.rgb=RGBColor.from_string('C64620')
    d.styles['Title'].font.size=Pt(26);d.styles['Heading 1'].font.size=Pt(20);d.styles['Heading 2'].font.size=Pt(12)
    d.add_paragraph('SYLLABUS TO SONG / '+DATE,'Subtitle')
    d.add_heading('Website + Android App',0)
    d.add_paragraph('We built both a responsive full-stack website and a native Android app for MBA Tech → CE → First year → Semester I. Both use the uploaded syllabus and reddish-orange sunrise design.')
    d.add_paragraph('2 deliverables • 91 topics • 10 ordered units • English and Hindi')
    p=d.add_paragraph('Webpage: ');hyperlink(p,WEB,WEB)
    d.add_paragraph('Website access is owner-private; sign in with the owning ChatGPT account if prompted.')
    p=d.add_paragraph('Android app: ');hyperlink(p,'Download the Android 8.0+ demonstration APK',APK)
    d.add_picture(str(SHOTS/'web-studio.png'),width=Inches(6.95))
    d.add_paragraph('Website dashboard — current UI with sample lyric drafts; no generated song is shown.','Caption')
    d.add_heading('Shared learning flow',2)
    d.add_paragraph('SAP login → programme/year/semester → Calculus or BEEE → chapter/topic → English or Hindi → music preferences → reviewed lyrics → music route → permitted playback/export.')
    d.add_paragraph('Website: responsive studio, notes, quizzes, syllabus report, library/player and reviewed source uploads. Android: native controls, local drafts/history/favourites, study tools, owner-GPU connection and Gemini browser handoff.')
    d.add_page_break()
    d.add_heading('Native Android companion',0)
    d.add_paragraph('Kotlin / Jetpack Compose • SQLite library • MediaPlayer • Android 8.0+')
    table=d.add_table(rows=2,cols=3)
    for i,(name,title) in enumerate([('android-studio.png','Chapter selection'),('android-music.png','Editable lyrics + ACE-Step action'),('android-notes.png','Source-linked study')]):
        table.cell(0,i).paragraphs[0].add_run().add_picture(str(SHOTS/name),width=Inches(2.14))
        table.cell(1,i).text=title
    d.add_paragraph('Actual app screens rendered in Robolectric; these are not physical-device screenshots. No audio was generated for the captures.','Caption')
    d.add_heading('Implemented features',2)
    for text in ['SAP sample login, all curriculum selectors and saved editable lyric drafts.','Genre, mood, voice, tempo, length, difficulty, formula repeats and custom instructions.','Notes, bilingual quizzes, topic completeness and demo-admin excerpt review.','History/favourites, foreground playback/seek/replay, permitted audio import/export.','GPU file import with Android Keystore encryption; Gemini browser fallback.']:
        d.add_paragraph(text,'List Bullet')
    d.add_heading('Actual audio requires a running music service',2)
    d.add_paragraph('More → Music connection → connect your running free ACE-Step GPU notebook → generate/edit lyrics → Synthesize with ACE-Step. Free GPU availability/session limits apply. The phone does not run the model.')
    d.add_page_break()
    d.add_heading('Features and final status',0)
    table=d.add_table(rows=4,cols=2)
    for i,(name,title) in enumerate([('web-notes.png','Notes & formula references'),('web-quiz.png','Quizzes with feedback'),('web-coverage.png','All-topic completeness report'),('web-gemini.png','Gemini / Lyria website handoff')]):
        row=(i//2)*2;col=i%2
        p=table.cell(row,col).paragraphs[0];iw,ih=Image.open(SHOTS/name).size;scale=min(3.2/iw,2.1/ih);p.add_run().add_picture(str(SHOTS/name),width=Inches(iw*scale),height=Inches(ih*scale))
        table.cell(row+1,col).text=title
    d.add_heading('Academic scope',2)
    d.add_paragraph('38 Calculus + 53 BEEE topics preserve the uploaded unit order, names, source summaries, formulas, locators and supplied objectives. Full chapter notes were not supplied. Explanations, drafts and quizzes are explicitly supplemental.')
    d.add_heading('Verification',2)
    d.add_paragraph('38 Python checks; 7 ACE protocol tests; 2 Robolectric/Compose flow tests; Android build/lint and APK signature verification. Website browser checks cover 91 report rows, Gemini navigation, isolated fixture playback and 320/390/768px layouts. GitHub website and Android workflows passed.')
    d.add_heading('Music-generation boundary',2)
    d.add_paragraph('No actual educational song or singing quality has been verified through this project. Gemini requires Create music → Paste → Send in its external page. ACE-Step needs an active GPU and a real generation/listening test. No song is fabricated. Native and website libraries are separate; physical-device audio decoding is unverified.')
    p=d.add_paragraph('Source: ');hyperlink(p,REPO,REPO)
    footer=sec.footer.paragraphs[0];footer.text='Syllabus to Song | Compact project report | '+DATE;footer.runs[0].font.size=Pt(8)
    d.save(str(OUT/'Syllabus_to_Song_Compact_Report.docx'))


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    pdf_report();word_report()
    (OUT/'capture-provenance.json').write_text(json.dumps({'report_date':DATE,'website':WEB,'apk':APK,
        'website_images':'Current code rendered in Chromium with isolated sample drafts; no generated/fixture audio seeded.',
        'android_images':'Current native Compose UI rendered using Robolectric SDK 35 native graphics; not physical-device screenshots.',
        'verified_audio_generation':False,'source_branch':'codex/development'},indent=2)+'\n')
    print('Built compact 3-page PDF and editable DOCX with real UI screenshots and clickable delivery links.')
if __name__=='__main__':main()
